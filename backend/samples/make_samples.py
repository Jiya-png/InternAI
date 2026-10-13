"""
Render the sample resumes in sample_texts.py to PDF / DOCX.

Run from ``backend``:   python -m samples.make_samples
(needs ``reportlab`` for PDFs:  pip install reportlab)
"""
from __future__ import annotations

from pathlib import Path

from samples.sample_texts import SAMPLES

OUT = Path(__file__).resolve().parent


def to_docx(text: str, path: Path) -> None:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    lines = text.strip().split("\n")
    d.add_heading(lines[0], level=1)
    for ln in lines[1:]:
        if ln.isupper() or (ln and len(ln.split()) <= 4 and ln[0].isupper() and not ln.startswith(("•", "-"))
                            and ln.istitle()):
            d.add_heading(ln, level=2)
        else:
            d.add_paragraph(ln)
    d.save(str(path))


def to_pdf(text: str, path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    font = "Helvetica"
    for cand in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf"):
        if Path(cand).exists():  # unicode font so ₹, •, – render
            pdfmetrics.registerFont(TTFont("Body", cand))
            font = "Body"
            break
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 50
    for i, ln in enumerate(text.strip().split("\n")):
        size = 16 if i == 0 else (11 if ln.isupper() else 9.5)
        c.setFont(font, size)
        # crude wrap at ~110 chars
        chunks = [ln[k:k + 110] for k in range(0, max(len(ln), 1), 110)] or [""]
        for ch in chunks:
            if y < 50:
                c.showPage()
                y = height - 50
                c.setFont(font, size)
            c.drawString(45, y, ch)
            y -= size + 4
    c.save()


def main() -> None:
    for name, (text, fmt) in SAMPLES.items():
        (OUT / f"{name}.txt").write_text(text, encoding="utf-8")
        path = OUT / f"{name}.{fmt}"
        (to_pdf if fmt == "pdf" else to_docx)(text, path)
        print("wrote", path)


if __name__ == "__main__":
    main()
